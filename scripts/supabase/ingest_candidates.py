"""Ingest external xls/xlsm question sources into public.question_candidates.

Parses each spreadsheet, normalizes the questions, auto-rejects rows that
already match an existing question in public.questions (`status='rejected'`,
`duplicate_of` set), and inserts the rest as `status='pending'` for review
in the dashboard. Idempotent: re-running won't insert the same (source_file,
source_nr) twice.

Run once locally:
    python3 scripts/supabase/ingest_candidates.py
"""
from __future__ import annotations

import json
import os
import re
import sys
import urllib.request
from pathlib import Path

import xlrd
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BASE = Path("/Users/donkfather/SynologyDrive/CAA Barca/OneDrive_1_10-05-2026/Chestionare/CHESTIONARE ANR CLASA C D")

SOURCES = [
    ("colreg",      "colreg agrement.xls"),
    ("maneuvering", "conducerea  manevra barcii.xls"),
    ("navigation",  "navigatie agrement.xls"),
    ("seamanship",  "marinarie agrement/- intrebari marinarie.xls"),
    ("rnd",         "RND agrement/intrebari RND agrementn.xlsm"),
]

# Layout per file: tuple (q_col, first_ans_col, stride, has_header)
# Pair files: question in col 1, answers start at col 2, then (ans, flag).
# Triple files: question in col 1, answers start at col 7, then (ans, url, flag).
# navigatie has no header at all and question lives in col 0.
LAYOUTS = {
    "colreg agrement.xls":                          {"q_col": 1, "ans_col": 2, "stride": 2, "has_header": True},
    "navigatie agrement.xls":                       {"q_col": 0, "ans_col": 1, "stride": 2, "has_header": False},
    "conducerea  manevra barcii.xls":               {"q_col": 1, "ans_col": 7, "stride": 3, "has_header": True},
    "marinarie agrement/- intrebari marinarie.xls": {"q_col": 1, "ans_col": 7, "stride": 3, "has_header": True},
    "RND agrement/intrebari RND agrementn.xlsm":    {"q_col": 1, "ans_col": 7, "stride": 3, "has_header": True},
}


def load_env() -> dict[str, str]:
    env = {}
    p = ROOT / ".env.local"
    for line in p.read_text().splitlines():
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


def read_sheet(path: Path) -> list[list]:
    if path.suffix == ".xls":
        wb = xlrd.open_workbook(str(path))
        sh = wb.sheet_by_index(0)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    wb = openpyxl.load_workbook(str(path), data_only=True)
    sh = wb[wb.sheetnames[0]]
    return [list(row) for row in sh.iter_rows(values_only=True)]


def parse_row(row: list, layout: dict) -> tuple[str, list[tuple[str, bool]], str | None] | None:
    """Return (question, [(answer, correct)], src_nr) or None if not a question row.

    Uses an explicit layout descriptor — each xls file has a different schema
    (some have section/type/difficulty metadata between the question and the
    first answer), and a naive next-non-empty-cell parser misreads them.
    """
    cells = [("" if v is None else str(v)).strip() for v in row]
    if not cells:
        return None

    q_col = layout["q_col"]
    ans_col = layout["ans_col"]
    stride = layout["stride"]

    nr = cells[0] if q_col > 0 and len(cells) > 0 and re.fullmatch(r"\d+(\.\d+)?", cells[0]) else None
    if nr:
        nr = nr.rstrip("0").rstrip(".") or cells[0]

    if q_col >= len(cells):
        return None
    qtext = cells[q_col]
    # Headers also live in column q_col on row 0 — skip them.
    if qtext.lower() in {"intrebare", "întrebare", ""}:
        return None
    if len(qtext) < 5:
        return None

    opts: list[tuple[str, bool]] = []
    i = ans_col
    while i + (stride - 1) < len(cells):
        ans = cells[i]
        flag = cells[i + stride - 1]  # flag is always the last cell in the group
        if ans and ans.lower() != "none":
            try:
                correct = bool(int(float(flag))) if flag != "" else False
            except (ValueError, TypeError):
                correct = False
            opts.append((ans, correct))
        i += stride
    if len(opts) < 2:
        return None
    return qtext, opts, nr


def normalize(s: str) -> str:
    s = re.sub(r"\s+", " ", s).strip().lower()
    s = (s.replace("ă", "a").replace("â", "a")
           .replace("î", "i").replace("ş", "s").replace("ș", "s")
           .replace("ţ", "t").replace("ț", "t"))
    s = re.sub(r"[?!.,;:\"'()\-–—]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def main() -> int:
    env = load_env()
    url = env["EXPO_PUBLIC_SUPABASE_URL"].rstrip("/")
    key = env["SUPABASE_SERVICE_KEY"]
    hdrs = {
        "Authorization": f"Bearer {key}",
        "apikey": key,
        "Content-Type": "application/json",
        "Prefer": "return=representation",
    }

    # Build dedup map from current public.questions
    req = urllib.request.Request(
        f"{url}/rest/v1/questions?select=id,question&order=id.asc",
        headers={"Authorization": f"Bearer {key}", "apikey": key},
    )
    with urllib.request.urlopen(req, timeout=30) as r:
        existing = json.loads(r.read())
    by_norm = {normalize(q["question"]): q["id"] for q in existing}

    # Build dedup map from already-ingested candidates (source_file, source_nr)
    req = urllib.request.Request(
        f"{url}/rest/v1/question_candidates?select=source_file,source_nr",
        headers={"Authorization": f"Bearer {key}", "apikey": key},
    )
    with urllib.request.urlopen(req, timeout=30) as r:
        already = json.loads(r.read())
    seen_keys = {(c["source_file"], c["source_nr"]) for c in already}

    new_rows: list[dict] = []
    counters = {"new": 0, "dup_db": 0, "dup_seen": 0, "skipped": 0}

    for topic, fname in SOURCES:
        path = BASE / fname
        if not path.exists():
            print(f"missing {fname}, skipping", file=sys.stderr)
            continue
        layout = LAYOUTS.get(fname)
        if not layout:
            print(f"no layout for {fname}, skipping", file=sys.stderr)
            continue
        rows = read_sheet(path)
        for r_idx, row in enumerate(rows):
            p = parse_row(row, layout)
            if not p:
                counters["skipped"] += 1
                continue
            qtext, opts, src_nr = p
            src_nr = src_nr or f"row{r_idx}"
            if (fname, src_nr) in seen_keys:
                counters["dup_seen"] += 1
                continue

            options = [a for a, _ in opts]
            try:
                correct = next(i for i, (_, c) in enumerate(opts) if c)
            except StopIteration:
                counters["skipped"] += 1
                continue

            norm = normalize(qtext)
            dup_id = by_norm.get(norm)

            cand = {
                "source_file": fname,
                "source_nr": src_nr,
                "topic": topic,
                "question": qtext,
                "options": options,
                "correct": correct,
                "license": [],
            }
            # Same keys on every row — PostgREST batch insert requires it.
            cand["status"] = "rejected" if dup_id is not None else "pending"
            cand["duplicate_of"] = dup_id
            if dup_id is not None:
                counters["dup_db"] += 1
            else:
                counters["new"] += 1
            new_rows.append(cand)

    print(f"to insert: {len(new_rows)}  ({counters})")

    # Batch insert
    if new_rows:
        for i in range(0, len(new_rows), 200):
            chunk = new_rows[i:i + 200]
            req = urllib.request.Request(
                f"{url}/rest/v1/question_candidates",
                data=json.dumps(chunk, ensure_ascii=False).encode("utf-8"),
                method="POST",
                headers=hdrs,
            )
            try:
                with urllib.request.urlopen(req, timeout=30) as r:
                    print(f"  inserted {len(chunk)} (HTTP {r.status})")
            except urllib.error.HTTPError as e:
                print(f"  ERROR {e.code}: {e.read().decode()[:500]}", file=sys.stderr)
                return 1

    return 0


if __name__ == "__main__":
    sys.exit(main())
